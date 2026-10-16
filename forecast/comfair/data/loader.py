"""
Benchmark data loader: reads official ComErgy 2026 Excel workbook.

Orchestrates extraction of:
- 8760 hourly demand periods (9 members)
- Tariffs (grid import, feedin)
- Environmental data (temperature, GHI)
- Member metadata (demographics, DGS assets)

Keeps workbook facts separate from modelling assumptions.
"""

import re
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import Optional, Tuple
import numpy as np
import openpyxl

from .models import BenchmarkData, MemberData, DGSInstallation


# Hardcoded indices into the official workbook structure
BENCHMARK_SHEET = "ComErgy Benchmark Dataset 2026"
HOURLY_DEMAND_COLS = list(range(1, 10))  # Columns A-I (1-indexed in openpyxl)
GRID_IMPORT_TARIFF_COL = 10  # Column J
FEEDIN_TARIFF_COL = 11  # Column K
TEMPERATURE_COL = 12  # Column L
GHI_COL = 13  # Column M
DATA_START_ROW = 3  # First numeric row (row 1=header, row 2=member labels, row 3=data)
DATA_END_ROW = 8762  # 8760 periods + header rows


def _parse_dgs_string(dgs_text: str) -> Tuple[int, int]:
    """
    Parse DGS installation string to extract PV and battery module counts.

    Examples:
        "15 PV modules" → (15, 0)
        "2 battery module" → (0, 2)
        "30 PV modules, 3 battery modules" → (30, 3)
        "Nothing" → (0, 0)

    Args:
        dgs_text: Raw text from workbook DGS column

    Returns:
        Tuple of (pv_modules, battery_modules)
    """
    if dgs_text.strip().lower() == "nothing":
        return 0, 0
    counts = {}
    for part in dgs_text.split(","):
        match = re.fullmatch(r"\s*(\d+)\s+(PV|battery)\s+modules?\s*", part, re.I)
        if not match or match[2].lower() in counts:
            raise ValueError(f"Invalid DGS installation: {dgs_text!r}")
        counts[match[2].lower()] = int(match[1])
    return counts.get("pv", 0), counts.get("battery", 0)


def _parse_income_string(income_text: str) -> float:
    """
    Parse annual income string to extract numeric value in AUD.

    Examples:
        "$200,000" → 200000.0
        "$80,000" → 80000.0
        "$70,000" → 70000.0

    Args:
        income_text: Raw text from workbook income column

    Returns:
        Annual income in AUD as float

    Raises:
        ValueError if income cannot be parsed
    """
    if not income_text:
        raise ValueError(f"Empty income text")

    # Remove $ and commas, convert to float
    cleaned = str(income_text).replace('$', '').replace(',', '').strip()
    try:
        return float(cleaned)
    except ValueError:
        raise ValueError(f"Cannot parse income '{income_text}' as float")


def _parse_occupant_string(occupant_text: str) -> int:
    """
    Parse occupant count string to extract numeric value.

    Examples:
        "2 people (adults)" → 2
        "5 people (2 adults, 3 children)" → 5
        "7 people (4 adults, 3 children)" → 7

    Args:
        occupant_text: Raw text from workbook occupant column

    Returns:
        Number of occupants as int

    Raises:
        ValueError if count cannot be parsed
    """
    if not occupant_text:
        raise ValueError(f"Empty occupant text")

    # Extract first number from string
    match = re.fullmatch(r'\s*(\d+) people \(.*\)\s*', str(occupant_text))
    if match:
        return int(match.group(1))

    raise ValueError(f"Cannot parse occupant count from '{occupant_text}'")


def load_benchmark(
    workbook_path: Optional[str] = None,
    validate: bool = True
) -> BenchmarkData:
    """
    Load official ComErgy 2026 benchmark dataset.

    Reads from Excel workbook, extracts 8760 hourly periods for 9 members,
    tariffs, weather, and member metadata. Returns typed BenchmarkData object.

    Args:
        workbook_path: Path to Excel file. If None, uses default:
            "backend/reader/ComErgy Benchmark Dataset 2026.xlsx"
        validate: If True, run BenchmarkData.validate() before returning

    Returns:
        BenchmarkData with all extracted fields populated

    Raises:
        FileNotFoundError: If workbook not found
        ValueError: If workbook structure is invalid or data is missing
    """
    if workbook_path is None:
        workbook_path = Path(__file__).resolve().parents[2] / "reader" / "ComErgy Benchmark Dataset 2026.xlsx"

    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(
            f"Benchmark workbook not found: {workbook_path.absolute()}"
        )

    # Load with openpyxl for raw access
    wb = openpyxl.load_workbook(str(workbook_path), data_only=False)

    try:
        if BENCHMARK_SHEET not in wb.sheetnames:
            raise ValueError(f"Expected sheet {BENCHMARK_SHEET!r} not found")
        return _load_sheet(wb[BENCHMARK_SHEET], workbook_path, validate)
    finally:
        wb.close()


def _load_sheet(ws, workbook_path: Path, validate: bool) -> BenchmarkData:
    """Workbook cell layout is confined to this ingestion boundary."""
    expected_headers = {
        "A1": "Hourly Electricity Demand (kW or kWh) -- 01 Jan-31 Dec",
        "J1": "Electricity price c/kWh", "K1": "Feedin Tariff c/kWh",
        "L1": "Dry Bulb Temperature [°C]", "M1": "GHI (W/m²)",
        "P2": "Location", "P3": "Price unit",
        "P4": "Daily connection fee to the grid",
        "P8": "Community members", "Q8": "DGS installation information",
        "R8": "Household type", "S8": "Number of occupants",
        "T8": "Annual household income (after tax)",
    }
    for coordinate, expected in expected_headers.items():
        if str(ws[coordinate].value).strip() != expected:
            raise ValueError(f"Unexpected header at {coordinate}: {ws[coordinate].value!r}")
    for member_id in range(1, 10):
        if ws.cell(2, member_id).value != f"Member {member_id}":
            raise ValueError("Demand member order must be Member 1 through Member 9")
    # Formatting extends below the dataset; only populated cells count as extra data.
    for row in ws.iter_rows(min_row=DATA_END_ROW + 1, max_col=13):
        if any(cell.value is not None for cell in row):
            raise ValueError("Expected exactly 8760 hourly periods; extra data found")
    for row in ws.iter_rows(min_row=18, min_col=16, max_col=16):
        if str(row[0].value).strip().startswith("Member "):
            raise ValueError("Unexpected extra member record")
    source_notes = {cell: str(ws[cell].value or "").strip()
                    for cell in ("A1", "J1", "K1", "Q3", "P19", "P20", "P22")}
    if not source_notes["P19"] or not source_notes["P20"]:
        raise ValueError("Missing module specification notes")
    # Extract hourly data (rows 3-8762, columns A-M)
    demand_list = []
    grid_import_list = []
    feedin_list = []
    temperature_list = []
    ghi_list = []

    for row_idx in range(DATA_START_ROW, DATA_END_ROW + 1):
        # Extract demand (columns A-I)
        row_demand = []
        for col_idx in HOURLY_DEMAND_COLS:
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is None or isinstance(cell_val, bool):
                raise ValueError(
                    f"Missing demand value at row {row_idx}, column {col_idx}"
                )
            try:
                row_demand.append(float(cell_val))
            except (ValueError, TypeError):
                raise ValueError(
                    f"Invalid demand value at row {row_idx}, column {col_idx}: "
                    f"{cell_val}"
                )
        demand_list.append(row_demand)

        # Extract tariff and environmental (columns J-M)
        for col_idx, col_list in [
            (GRID_IMPORT_TARIFF_COL, grid_import_list),
            (FEEDIN_TARIFF_COL, feedin_list),
            (TEMPERATURE_COL, temperature_list),
            (GHI_COL, ghi_list),
        ]:
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is None or isinstance(cell_val, bool):
                raise ValueError(
                    f"Missing value at row {row_idx}, column {col_idx}"
                )
            try:
                col_list.append(float(cell_val))
            except (ValueError, TypeError):
                raise ValueError(
                    f"Invalid value at row {row_idx}, column {col_idx}: {cell_val}"
                )

    # Convert to numpy arrays
    demand_hourly = np.array(demand_list, dtype=np.float64)
    grid_import_tariff = np.array(grid_import_list, dtype=np.float64)
    feedin_tariff = np.array(feedin_list, dtype=np.float64)
    temperature_c = np.array(temperature_list, dtype=np.float64)
    ghi_w_per_m2 = np.array(ghi_list, dtype=np.float64)

    # Verify period count
    if len(demand_hourly) != 8760:
        raise ValueError(
            f"Expected 8760 hourly periods, got {len(demand_hourly)}"
        )

    # Extract static metadata (from row 2-4 area, columns P-Q)
    location = None
    location_label = ws.cell(row=2, column=16).value  # P2 (col 16 = P)
    if location_label and "location" in str(location_label).lower():
        location = ws.cell(row=2, column=17).value  # Q2 (col 17 = Q)
        if location:
            location = str(location).strip()

    price_unit = None
    price_unit_label = ws.cell(row=3, column=16).value  # P3 (col 16 = P)
    if price_unit_label and "price" in str(price_unit_label).lower():
        price_unit_val = ws.cell(row=3, column=17).value  # Q3 (col 17 = Q)
        if price_unit_val:
            # Extract "AUD" or "Australian Dollar"
            unit_str = str(price_unit_val).strip()
            if "AUD" in unit_str or "Australian" in unit_str:
                price_unit = "AUD"

    daily_connection_fee = None
    connection_fee_label = ws.cell(row=4, column=16).value  # P4 (col 16 = P)
    if connection_fee_label and "connection" in str(connection_fee_label).lower():
        connection_fee_val = ws.cell(row=4, column=17).value  # Q4 (col 17 = Q)
        if connection_fee_val is not None:
            try:
                daily_connection_fee = float(connection_fee_val)
            except (ValueError, TypeError):
                raise ValueError(
                    f"Invalid connection fee value: {connection_fee_val}"
                )

    if location is None:
        raise ValueError("Could not extract location from workbook")
    if price_unit is None:
        raise ValueError("Could not extract price unit from workbook")
    if daily_connection_fee is None:
        raise ValueError("Could not extract daily connection fee from workbook")

    # Extract member metadata (rows 9-17, columns P-T)
    members = {}
    for member_id in range(1, 10):
        row_idx = 8 + member_id  # Row 9 for Member 1, row 10 for Member 2, etc.

        # Member name (P column, col 16)
        member_name = ws.cell(row=row_idx, column=16).value
        if not member_name or str(member_name).strip() != f"Member {member_id}":
            raise ValueError(
                f"Expected 'Member {member_id}' at row {row_idx}, "
                f"got '{member_name}'"
            )

        # DGS installation info (Q column, col 17)
        dgs_text = ws.cell(row=row_idx, column=17).value
        pv_modules, battery_modules = _parse_dgs_string(str(dgs_text))

        # Household type (R column, col 18)
        household_type = ws.cell(row=row_idx, column=18).value or ""
        household_type = str(household_type).strip()

        # Occupant count (S column, col 19)
        occupant_text = ws.cell(row=row_idx, column=19).value
        occupant_count = _parse_occupant_string(str(occupant_text))

        # Annual income (T column, col 20)
        income_text = ws.cell(row=row_idx, column=20).value
        annual_income_aud = _parse_income_string(str(income_text))

        # Create member object
        dgs = DGSInstallation(
            member_id=member_id,
            pv_modules=pv_modules,
            battery_modules=battery_modules,
        )
        member = MemberData(
            member_id=member_id,
            household_type=household_type,
            occupant_count=occupant_count,
            annual_income_aud=annual_income_aud,
            dgs=dgs,
        )
        members[member_id] = member

    # Create BenchmarkData object
    benchmark = BenchmarkData(
        workbook_path=str(workbook_path),
        workbook_sheet=BENCHMARK_SHEET,
        load_timestamp=datetime.now(timezone.utc),
        source_notes=source_notes,
        workbook_sha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        member_count=9,
        period_count=8760,
        demand_hourly=demand_hourly,
        grid_import_tariff_raw=grid_import_tariff,
        feedin_tariff_raw=feedin_tariff,
        temperature_c=temperature_c,
        ghi_w_per_m2=ghi_w_per_m2,
        location=location,
        price_unit=price_unit,
        daily_connection_fee=daily_connection_fee,
        daily_connection_fee_unit=price_unit,  # Currency note Q3; fee is daily per P4
        members=members,
    )

    if validate:
        benchmark.validate()

    return benchmark
